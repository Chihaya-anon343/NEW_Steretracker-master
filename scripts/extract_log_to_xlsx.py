#!/usr/bin/env python3
"""从 tracking_log.txt 提取每帧位姿结果写入 xlsx。

提取字段: 帧序数 / rvec(xyz) / tvec(xyz, mm) / 择优重投影误差(px) / 状态。
PnP 失败帧保留行, 数值留空, 状态=失败。

用法:
    python scripts/extract_log_to_xlsx.py [log_path] [xlsx_path]
默认:
    log_path  = output/image_00/tracking_log.txt
    xlsx_path = <log 同目录>/tracking_results.xlsx
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FRAME_RE = re.compile(r"^===== 第 (\d+) 帧")
POSE_RE = re.compile(r"Pose: rvec=\[([^\]]+)\]\s+tvec=\[([^\]]+)\]")
REPROJ_RE = re.compile(r"择优: 重投影=([-+0-9.eE]+)px")

HEADERS = ["帧序数", "rvec_x", "rvec_y", "rvec_z",
           "tvec_x (mm)", "tvec_y (mm)", "tvec_z (mm)",
           "重投影误差 (px)", "状态"]


def parse_log(log_path):
    rows = []
    cur = None

    def new_row(frame_no):
        return {"frame": frame_no, "rvec": None, "tvec": None, "reproj": None}

    with open(log_path, "r", encoding="utf-8", errors="replace") as f:
        for line in f:
            m = FRAME_RE.match(line)
            if m:
                if cur is not None:
                    rows.append(cur)
                cur = new_row(int(m.group(1)))
                continue
            if cur is None:
                continue
            if cur["rvec"] is None:
                m = POSE_RE.search(line)
                if m:
                    cur["rvec"] = [float(v) for v in m.group(1).split(",")]
                    cur["tvec"] = [float(v) for v in m.group(2).split(",")]
                    continue
            if cur["reproj"] is None:
                m = REPROJ_RE.search(line)
                if m:
                    cur["reproj"] = float(m.group(1))
        if cur is not None:
            rows.append(cur)
    return rows


def write_xlsx(rows, xlsx_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "tracking"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    fail_fill = PatternFill("solid", fgColor="FCE4EC")
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill

    for i, r in enumerate(rows, 2):
        ok = r["rvec"] is not None and r["tvec"] is not None
        values = [r["frame"]]
        if ok:
            values += r["rvec"] + r["tvec"]
        else:
            values += [None] * 6
        values += [r["reproj"] if r["reproj"] is not None else None,
                   "成功" if ok else "失败"]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=i, column=col, value=v)
            if not ok:
                c.fill = fail_fill

    widths = [10, 14, 14, 14, 14, 14, 14, 16, 8]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"

    wb.save(xlsx_path)


def main():
    log_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("output/image_00/tracking_log.txt")
    if len(sys.argv) > 2:
        xlsx_path = Path(sys.argv[2])
    else:
        xlsx_path = log_path.parent / "tracking_results.xlsx"

    rows = parse_log(log_path)
    n_ok = sum(1 for r in rows if r["rvec"] is not None)
    write_xlsx(rows, xlsx_path)
    print(f"帧总数: {len(rows)}  成功: {n_ok}  失败: {len(rows) - n_ok}")
    print(f"已写入: {xlsx_path}")


if __name__ == "__main__":
    main()
